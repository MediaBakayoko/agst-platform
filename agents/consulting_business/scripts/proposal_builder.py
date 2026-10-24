#!/usr/bin/env python3
"""
Genere un brouillon d'offre technique (Word) a partir d'une analyse de
ToR (JSON produit par tor_analyzer.py) et d'un profil consultant/cabinet
reel (JSON fourni par l'utilisateur).

Ne mentionne JAMAIS une experience, reference ou qualification absente
du profil fourni -- si le profil est incomplet, le dit explicitement
dans le document plutot que d'inventer.

Usage:
    python proposal_builder.py \
        --tor-analysis output/tor_test_XXX/tor_analysis.json \
        --profile sample_data/profil_consultant.json \
        --output-dir output/
"""
from __future__ import annotations

import argparse
import json
import logging
import sys
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.shared import Mm, RGBColor

logger = logging.getLogger("consulting_business.proposal_builder")


def build_proposal(tor_analysis: dict, profile: dict, output_path: Path) -> list[str]:
    """Construit le document Word. Renvoie la liste des avertissements
    (sections ToR non trouvees, elements de profil absents) a signaler
    a l'utilisateur.
    """
    warnings = []
    doc = Document()
    section = doc.sections[0]
    section.left_margin = Mm(20)
    section.right_margin = Mm(20)

    title = doc.add_heading("Proposition technique (brouillon)", level=0)
    for run in title.runs:
        run.font.color.rgb = RGBColor(0x1F, 0x38, 0x64)

    doc.add_paragraph(
        "⚠ BROUILLON — document non finalise. A relire et completer avant tout envoi. "
        "Aucune donnee financiere ou reference n'a ete inventee ; les elements manquants "
        "sont signales ci-dessous."
    ).italic = True

    doc.add_heading("1. Rappel des termes de reference", level=1)
    sections = tor_analysis.get("sections", {})
    labels = {
        "objectifs": "Objectifs de la mission",
        "livrables": "Livrables attendus",
        "echeance": "Échéance / durée",
        "budget_plafond": "Budget / enveloppe (si mentionné dans le ToR)",
        "qualifications_requises": "Qualifications requises",
    }
    for key, label in labels.items():
        doc.add_heading(label, level=2)
        content = sections.get(key)
        if content:
            doc.add_paragraph(content)
        else:
            doc.add_paragraph(
                "Non trouvé dans le ToR fourni — à compléter manuellement après relecture du document original."
            ).italic = True
            warnings.append(f"Section ToR non trouvee : {label}")

    doc.add_heading("2. Présentation du prestataire", level=1)
    nom = profile.get("nom_structure") or profile.get("nom_consultant")
    if not nom:
        doc.add_paragraph("Nom du prestataire non renseigné dans le profil fourni.").italic = True
        warnings.append("Profil : nom du prestataire absent")
    else:
        doc.add_paragraph(nom)

    doc.add_heading("Références / expériences pertinentes", level=2)
    experiences = profile.get("experiences", [])
    if not experiences:
        doc.add_paragraph(
            "Aucune expérience renseignée dans le profil fourni — à compléter avec des "
            "références réelles avant envoi. Ne pas inventer de référence client."
        ).italic = True
        warnings.append("Profil : aucune expérience/référence renseignée")
    else:
        for exp in experiences:
            doc.add_paragraph(f"• {exp}", style="List Bullet")

    doc.add_heading("Qualifications / certifications", level=2)
    quals = profile.get("qualifications", [])
    if not quals:
        doc.add_paragraph("Aucune qualification renseignée dans le profil fourni.").italic = True
        warnings.append("Profil : aucune qualification renseignée")
    else:
        for q in quals:
            doc.add_paragraph(f"• {q}", style="List Bullet")

    doc.add_heading("3. Approche méthodologique proposée", level=1)
    doc.add_paragraph(
        "[À RÉDIGER] — section à compléter manuellement : démarche technique proposée "
        "en réponse aux objectifs et livrables ci-dessus."
    ).italic = True
    warnings.append("Section methodologique non redigee automatiquement -- a completer manuellement")

    doc.add_heading("4. Budget", level=1)
    doc.add_paragraph(
        "[VOIR FICHIER BUDGET SEPARE] — généré par budget_builder.py. "
        "Tout montant marqué [HYPOTHESE] n'est pas encore validé par le client."
    ).italic = True

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)
    return warnings


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--tor-analysis", required=True)
    ap.add_argument("--profile", required=True)
    ap.add_argument("--output-dir", default="output")
    args = ap.parse_args()

    tor_path = Path(args.tor_analysis)
    profile_path = Path(args.profile)
    if not tor_path.exists():
        print(f"Analyse ToR introuvable : {tor_path}", file=sys.stderr)
        return 1
    if not profile_path.exists():
        print(f"Profil introuvable : {profile_path}", file=sys.stderr)
        return 1

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    run_dir = Path(args.output_dir) / f"proposition_{timestamp}"
    run_dir.mkdir(parents=True, exist_ok=True)
    logging.basicConfig(level=logging.INFO, format="%(asctime)s [%(levelname)s] %(message)s",
                         handlers=[logging.FileHandler(run_dir / "journal.log", encoding="utf-8"), logging.StreamHandler()])

    tor_analysis = json.loads(tor_path.read_text(encoding="utf-8"))
    profile = json.loads(profile_path.read_text(encoding="utf-8"))

    output_path = run_dir / "proposition_technique_brouillon.docx"
    warnings = build_proposal(tor_analysis, profile, output_path)

    for w in warnings:
        logger.warning(w)
    logger.info("Proposition ecrite : %s (%d avertissement(s))", output_path, len(warnings))

    print(f"OK - proposition: {output_path}")
    if warnings:
        print(f"ATTENTION : {len(warnings)} element(s) a completer manuellement avant envoi")
    print("RAPPEL : ceci est un brouillon -- aucun envoi automatique, validation humaine requise.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
