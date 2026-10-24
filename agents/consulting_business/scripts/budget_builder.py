#!/usr/bin/env python3
"""
Construit un budget detaille (Excel) a partir de lignes fournies par
l'utilisateur (CSV : poste, quantite, unite, taux_journalier_ou_unitaire,
hypothese).

Ne propose JAMAIS de taux par defaut invente -- toute ligne sans taux
fourni est signalee en erreur, pas completee automatiquement.

CSV attendu (colonnes) :
    poste, quantite, unite, taux, hypothese
    hypothese = "oui" si le taux n'est pas encore valide par le client,
                "non" si c'est un montant deja accepte/contractuel.

Usage:
    python budget_builder.py sample_data/lignes_budget_test.csv --output-dir output/
"""
from __future__ import annotations

import argparse
import logging
import sys
from datetime import datetime
from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

logger = logging.getLogger("consulting_business.budget_builder")


def build_budget(df: pd.DataFrame) -> tuple[pd.DataFrame, list[str]]:
    errors = []
    required_cols = {"poste", "quantite", "unite", "taux", "hypothese"}
    missing_cols = required_cols - set(df.columns)
    if missing_cols:
        errors.append(f"Colonnes manquantes dans le CSV : {missing_cols}")
        return df, errors

    df = df.copy()
    for idx, row in df.iterrows():
        if pd.isna(row["taux"]) or row["taux"] == "":
            errors.append(f"Ligne {idx} ('{row['poste']}') : aucun taux fourni -- impossible de calculer, montant non estime")

    df["quantite"] = pd.to_numeric(df["quantite"], errors="coerce")
    df["taux"] = pd.to_numeric(df["taux"], errors="coerce")
    df["montant"] = df["quantite"] * df["taux"]
    df["statut"] = df["hypothese"].apply(lambda h: "[HYPOTHESE - non valide]" if str(h).strip().lower() in ("oui", "yes", "true", "1") else "[valide/contractuel]")
    return df, errors


def write_excel(df: pd.DataFrame, errors: list[str], output_path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Budget"
    header_fill = PatternFill(start_color="1F3864", end_color="1F3864", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)

    headers = ["poste", "quantite", "unite", "taux", "montant", "statut"]
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.fill = header_fill
        cell.font = header_font

    for i, row in enumerate(df.itertuples(index=False), start=2):
        row_dict = dict(zip(df.columns, row))
        for c, h in enumerate(headers, start=1):
            ws.cell(row=i, column=c, value=row_dict.get(h))
        if "HYPOTHESE" in str(row_dict.get("statut", "")):
            fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
            for c in range(1, len(headers) + 1):
                ws.cell(row=i, column=c).fill = fill

    total_row = len(df) + 3
    ws.cell(row=total_row, column=1, value="TOTAL").font = Font(bold=True)
    ws.cell(row=total_row, column=5, value=f"=SUM(E2:E{len(df)+1})").font = Font(bold=True)

    for c in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(c)].width = 22

    if errors:
        ws2 = wb.create_sheet("Erreurs")
        ws2.cell(row=1, column=1, value="Lignes sans taux fourni -- montant NON estime automatiquement").font = Font(bold=True, color="C00000")
        for i, err in enumerate(errors, start=2):
            ws2.cell(row=i, column=1, value=err)
        ws2.column_dimensions["A"].width = 90

    wb.save(output_path)


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("budget_csv")
    ap.add_argument("--output-dir", default="output")
    args = ap.parse_args()

    csv_path = Path(args.budget_csv)
    if not csv_path.exists():
        print(f"Fichier introuvable : {csv_path}", file=sys.stderr)
        return 1

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    run_dir = Path(args.output_dir) / f"{csv_path.stem}_{timestamp}"
    run_dir.mkdir(parents=True, exist_ok=True)
    logging.basicConfig(level=logging.INFO, format="%(asctime)s [%(levelname)s] %(message)s",
                         handlers=[logging.FileHandler(run_dir / "journal.log", encoding="utf-8"), logging.StreamHandler()])

    df_raw = pd.read_csv(csv_path)
    df_budget, errors = build_budget(df_raw)
    for e in errors:
        logger.warning(e)

    excel_path = run_dir / "budget.xlsx"
    write_excel(df_budget, errors, excel_path)
    logger.info("Budget ecrit : %s (%d ligne(s), %d erreur(s))", excel_path, len(df_budget), len(errors))

    print(f"OK - budget: {excel_path}")
    if errors:
        print(f"ATTENTION : {len(errors)} ligne(s) sans taux fourni -- voir feuille 'Erreurs'")
    return 0


if __name__ == "__main__":
    sys.exit(main())
