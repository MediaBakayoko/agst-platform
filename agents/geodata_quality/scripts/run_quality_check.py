#!/usr/bin/env python3
"""
CLI du moteur GeoData Quality.

Usage:
    python run_quality_check.py <input_file> \
        --output-dir output/ \
        --required-fields nom,surface_ha \
        --numeric-fields surface_ha \
        --area-field surface_ha \
        --area-tolerance-pct 20 \
        --dup-tolerance-m 0 \
        --id-field id_parcelle

Ne modifie jamais le fichier d'entree. Toute sortie va dans un
sous-dossier horodate de --output-dir.
"""
from __future__ import annotations

import argparse
import logging
import sys
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).parent))
from quality_checks import run_quality_checks  # noqa: E402


def _parse_list(value: str | None) -> list[str]:
    if not value:
        return []
    return [v.strip() for v in value.split(",") if v.strip()]


def write_excel_report(result, output_path: Path) -> None:
    wb = openpyxl.Workbook()

    # --- Feuille Resume ---
    ws = wb.active
    ws.title = "Resume"
    header_fill = PatternFill(start_color="1F3864", end_color="1F3864", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)

    rows = [
        ("Fichier source", result.input_path),
        ("Nombre d'entites", result.n_features),
        ("Entites avec au moins une anomalie", len(result.anomaly_feature_ids)),
        ("CRS d'origine", result.crs_origine or "(non defini)"),
        ("CRS utilise pour calculs de surface", result.crs_calcul or "(non calcule)"),
        ("Debut du traitement (UTC)", result.started_at),
        ("Fin du traitement (UTC)", result.finished_at),
        ("Nombre total d'anomalies", len(result.anomalies)),
    ]
    for r, (label, value) in enumerate(rows, start=1):
        ws.cell(row=r, column=1, value=label).font = Font(bold=True)
        ws.cell(row=r, column=2, value=value)

    ws.cell(row=len(rows) + 2, column=1, value="Repartition par type d'anomalie").font = Font(bold=True, size=12)
    summary_df = result.summary_by_type()
    start_row = len(rows) + 3
    headers = ["type_anomalie", "severite", "nombre"]
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=start_row, column=c, value=h)
        cell.fill = header_fill
        cell.font = header_font
    for i, row in enumerate(summary_df.itertuples(index=False), start=1):
        for c, val in enumerate(row, start=1):
            ws.cell(row=start_row + i, column=c, value=val)

    for col in range(1, 4):
        ws.column_dimensions[get_column_letter(col)].width = 35

    # --- Feuille Anomalies detaillees ---
    ws2 = wb.create_sheet("Anomalies")
    detail_df = result.to_dataframe()
    headers2 = list(detail_df.columns)
    for c, h in enumerate(headers2, start=1):
        cell = ws2.cell(row=1, column=c, value=h)
        cell.fill = header_fill
        cell.font = header_font
    severity_fill = {
        "erreur": PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid"),
        "a_verifier": PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid"),
    }
    for i, row in enumerate(detail_df.itertuples(index=False), start=2):
        row_dict = dict(zip(headers2, row))
        for c, h in enumerate(headers2, start=1):
            ws2.cell(row=i, column=c, value=row_dict[h])
        sev = row_dict.get("severite")
        if sev in severity_fill:
            for c in range(1, len(headers2) + 1):
                ws2.cell(row=i, column=c).fill = severity_fill[sev]
    for col in range(1, len(headers2) + 1):
        ws2.column_dimensions[get_column_letter(col)].width = 28

    # --- Feuille Parametres ---
    ws3 = wb.create_sheet("Parametres")
    ws3.cell(row=1, column=1, value="Parametre").font = Font(bold=True)
    ws3.cell(row=1, column=2, value="Valeur").font = Font(bold=True)
    for i, (k, v) in enumerate(result.params.items(), start=2):
        ws3.cell(row=i, column=1, value=k)
        ws3.cell(row=i, column=2, value=str(v))
    ws3.column_dimensions["A"].width = 30
    ws3.column_dimensions["B"].width = 50

    wb.save(output_path)


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("input_file", help="Fichier geospatial a controler (GeoPackage, Shapefile, GeoJSON, CSV geolocalise)")
    ap.add_argument("--output-dir", default="output", help="Dossier racine des sorties (defaut: output/)")
    ap.add_argument("--layer", default=None, help="Nom de la couche si le fichier en contient plusieurs (GeoPackage)")
    ap.add_argument("--id-field", default=None, help="Champ a utiliser comme identifiant d'entite dans le rapport")
    ap.add_argument("--required-fields", default="", help="Champs obligatoires, separes par des virgules")
    ap.add_argument("--numeric-fields", default="", help="Champs qui doivent etre numeriques, separes par des virgules")
    ap.add_argument("--dedup-key-fields", default="", help="Champs cles pour detecter les doublons attributaires")
    ap.add_argument("--area-field", default=None, help="Champ de surface declaree (en hectares) a comparer au calcul SIG")
    ap.add_argument("--area-tolerance-pct", type=float, default=20.0, help="Tolerance d'ecart surface declaree vs SIG, en %% (defaut 20)")
    ap.add_argument("--dup-tolerance-m", type=float, default=0.0, help="Tolerance de distance pour les doublons geometriques, en metres (defaut 0 = identite stricte)")
    args = ap.parse_args()

    input_path = Path(args.input_file)
    if not input_path.exists():
        print(f"Fichier introuvable : {input_path}", file=sys.stderr)
        return 1

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    run_dir = Path(args.output_dir) / f"{input_path.stem}_{timestamp}"
    run_dir.mkdir(parents=True, exist_ok=True)

    log_path = run_dir / "journal.log"
    logging.basicConfig(
        level=logging.INFO,
        format="%(asctime)s [%(levelname)s] %(message)s",
        handlers=[logging.FileHandler(log_path, encoding="utf-8"), logging.StreamHandler()],
    )
    logger = logging.getLogger("geodata_quality.cli")
    logger.info("Demarrage du controle qualite sur %s", input_path)
    logger.info("Parametres CLI : %s", vars(args))

    result = run_quality_checks(
        input_path=str(input_path),
        required_fields=_parse_list(args.required_fields),
        numeric_fields=_parse_list(args.numeric_fields),
        key_fields_for_dedup=_parse_list(args.dedup_key_fields),
        area_field=args.area_field,
        area_tolerance_pct=args.area_tolerance_pct,
        dup_tolerance_m=args.dup_tolerance_m,
        layer=args.layer,
        id_field=args.id_field,
    )

    excel_path = run_dir / "rapport_qualite.xlsx"
    write_excel_report(result, excel_path)
    logger.info("Rapport Excel ecrit : %s", excel_path)

    gdf = getattr(result, "_gdf", None)
    gpkg_path = None
    if gdf is not None and result.anomaly_feature_ids:
        anomalous = gdf[gdf["_feature_id"].isin(result.anomaly_feature_ids)].copy()
        if not anomalous.empty:
            gpkg_path = run_dir / "anomalies.gpkg"
            anomalous.drop(columns=["_feature_id"], errors="ignore").to_file(gpkg_path, driver="GPKG")
            logger.info("Couche des anomalies ecrite : %s (%d entites)", gpkg_path, len(anomalous))

    logger.info(
        "Termine. %d entite(s) analysee(s), %d anomalie(s) detectee(s) sur %d entite(s) concernee(s).",
        result.n_features, len(result.anomalies), len(result.anomaly_feature_ids),
    )

    print(f"OK - rapport: {excel_path}")
    if gpkg_path:
        print(f"OK - couche anomalies: {gpkg_path}")
    print(f"OK - journal: {log_path}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
