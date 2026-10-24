#!/usr/bin/env python3
"""
CLI du pipeline Field2Report : collecte -> controle qualite -> nettoyage
-> analyse -> cartographie -> rapport.

Usage:
    python run_pipeline.py <input_csv_or_xlsx> \
        --output-dir output/ \
        --lat-field latitude --lon-field longitude \
        --group-field commune \
        --numeric-fields age,surface_ha,rendement_kg_ha \
        --required-fields nom_repondant,commune \
        --text-fields nom_repondant,commune \
        --dedup-key-fields nom_repondant,commune
"""
from __future__ import annotations

import argparse
import logging
import sys
from datetime import datetime
from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).parent))
from cleaning import run_cleaning_pipeline  # noqa: E402
from analysis import descriptive_stats, detect_sample_bias, plot_survey_map  # noqa: E402
from report_builder import build_report  # noqa: E402


def _parse_list(value: str | None) -> list[str]:
    if not value:
        return []
    return [v.strip() for v in value.split(",") if v.strip()]


def write_stats_excel(stats_df: pd.DataFrame, anomalies: list[dict], output_path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Statistiques"
    header_fill = PatternFill(start_color="1F3864", end_color="1F3864", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)

    cols = list(stats_df.columns) if not stats_df.empty else ["variable", "groupe", "n"]
    for c, h in enumerate(cols, start=1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.fill = header_fill
        cell.font = header_font
    for i, row in enumerate(stats_df.itertuples(index=False), start=2):
        for c, val in enumerate(row, start=1):
            ws.cell(row=i, column=c, value=val)
    for c in range(1, len(cols) + 1):
        ws.column_dimensions[get_column_letter(c)].width = 18

    ws2 = wb.create_sheet("Anomalies")
    headers2 = ["type_anomalie", "severite", "ligne", "champ", "description"]
    for c, h in enumerate(headers2, start=1):
        cell = ws2.cell(row=1, column=c, value=h)
        cell.fill = header_fill
        cell.font = header_font
    for i, a in enumerate(anomalies, start=2):
        for c, h in enumerate(headers2, start=1):
            ws2.cell(row=i, column=c, value=a.get(h))
    for c in range(1, len(headers2) + 1):
        ws2.column_dimensions[get_column_letter(c)].width = 25

    wb.save(output_path)


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("input_file", help="CSV ou Excel issu de la collecte (KoboToolbox/ODK export, etc.)")
    ap.add_argument("--output-dir", default="output")
    ap.add_argument("--lat-field", default=None)
    ap.add_argument("--lon-field", default=None)
    ap.add_argument("--group-field", default=None, help="Champ de regroupement (ex: commune, district)")
    ap.add_argument("--numeric-fields", default="")
    ap.add_argument("--text-fields", default="")
    ap.add_argument("--required-fields", default="")
    ap.add_argument("--dedup-key-fields", default="")
    args = ap.parse_args()

    input_path = Path(args.input_file)
    if not input_path.exists():
        print(f"Fichier introuvable : {input_path}", file=sys.stderr)
        return 1

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    run_dir = Path(args.output_dir) / f"{input_path.stem}_{timestamp}"
    run_dir.mkdir(parents=True, exist_ok=True)

    log_path = run_dir / "journal.log"
    logging.basicConfig(
        level=logging.INFO,
        format="%(asctime)s [%(levelname)s] %(message)s",
        handlers=[logging.FileHandler(log_path, encoding="utf-8"), logging.StreamHandler()],
    )
    logger = logging.getLogger("field2report.cli")
    started_at = datetime.utcnow().isoformat() + "Z"
    logger.info("Demarrage du pipeline sur %s", input_path)
    logger.info("Parametres CLI : %s", vars(args))

    if input_path.suffix.lower() in (".xlsx", ".xls"):
        df_raw = pd.read_excel(input_path)
    else:
        df_raw = pd.read_csv(input_path)
    logger.info("Donnees brutes chargees : %d lignes, %d colonnes", len(df_raw), len(df_raw.columns))

    numeric_fields = _parse_list(args.numeric_fields)
    text_fields = _parse_list(args.text_fields)
    required_fields = _parse_list(args.required_fields)
    dedup_key_fields = _parse_list(args.dedup_key_fields)

    result = run_cleaning_pipeline(
        df_raw,
        required_fields=required_fields,
        numeric_fields=numeric_fields,
        text_fields=text_fields,
        lat_field=args.lat_field,
        lon_field=args.lon_field,
        dedup_key_fields=dedup_key_fields,
    )
    logger.info("Nettoyage termine : %d action(s), %d anomalie(s) detectee(s)",
                len(result.actions), len(result.anomalies))
    for a in result.actions:
        logger.info("  action: %s champ=%s lignes=%d", a.action, a.champ, a.n_lignes_concernees)

    clean_path = run_dir / "cleaned.csv"
    result.df_clean.to_csv(clean_path, index=False)
    logger.info("Donnees nettoyees ecrites : %s", clean_path)

    stats_df = descriptive_stats(result.df_clean, numeric_fields, group_field=args.group_field)
    stats_excel_path = run_dir / "tableaux_statistiques.xlsx"
    write_stats_excel(stats_df, result.anomalies, stats_excel_path)
    logger.info("Tableaux statistiques ecrits : %s", stats_excel_path)

    bias_notes = detect_sample_bias(result.df_clean, group_field=args.group_field)
    for note in bias_notes:
        logger.info("Biais signale : %s", note)

    map_path = None
    if args.lat_field and args.lon_field:
        map_path = plot_survey_map(
            result.df_clean, args.lat_field, args.lon_field,
            run_dir / "carte_enquete.png", group_field=args.group_field,
        )
        if map_path:
            logger.info("Carte generee : %s", map_path)

    finished_at = datetime.utcnow().isoformat() + "Z"
    report_path = run_dir / "rapport_final.docx"
    build_report(
        output_path=report_path,
        source_file=str(input_path),
        n_rows_raw=len(df_raw),
        n_rows_clean=len(result.df_clean),
        cleaning_actions=result.actions,
        anomalies=result.anomalies,
        stats_df=stats_df,
        bias_notes=bias_notes,
        map_path=map_path,
        group_field=args.group_field,
        started_at=started_at,
        finished_at=finished_at,
    )
    logger.info("Rapport Word ecrit : %s", report_path)

    print(f"OK - donnees nettoyees: {clean_path}")
    print(f"OK - statistiques: {stats_excel_path}")
    if map_path:
        print(f"OK - carte: {map_path}")
    print(f"OK - rapport: {report_path}")
    print(f"OK - journal: {log_path}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
