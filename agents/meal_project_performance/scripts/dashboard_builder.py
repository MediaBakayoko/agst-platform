#!/usr/bin/env python3
"""
Construit un tableau de bord Excel (indicateurs x periodes) a partir
du fichier produit par indicator_calculator.py, avec code couleur par
statut et graphique de progression.

Usage:
    python dashboard_builder.py --calculated output/XXX/indicateurs_calcules.csv --output-dir output/
"""
from __future__ import annotations

import argparse
import logging
import sys
from datetime import datetime
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import openpyxl
import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

logger = logging.getLogger("meal.dashboard_builder")

STATUS_COLORS = {
    "atteint": "C6EFCE",
    "en_bonne_voie": "FFEB9C",
    "en_retard": "FFC966",
    "tres_en_retard": "FFC7CE",
    "non_calculable": "D9D9D9",
}


def write_dashboard(df: pd.DataFrame, output_path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Tableau de bord"
    header_fill = PatternFill(start_color="1F3864", end_color="1F3864", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)

    headers = list(df.columns)
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.fill = header_fill
        cell.font = header_font

    for i, row in enumerate(df.itertuples(index=False), start=2):
        row_dict = dict(zip(headers, row))
        for c, h in enumerate(headers, start=1):
            ws.cell(row=i, column=c, value=row_dict.get(h))
        statut = row_dict.get("statut")
        if statut in STATUS_COLORS:
            fill = PatternFill(start_color=STATUS_COLORS[statut], end_color=STATUS_COLORS[statut], fill_type="solid")
            for c in range(1, len(headers) + 1):
                ws.cell(row=i, column=c).fill = fill

    for c in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(c)].width = 20

    ws2 = wb.create_sheet("Synthese par statut")
    counts = df["statut"].value_counts()
    ws2.cell(row=1, column=1, value="Statut").font = Font(bold=True)
    ws2.cell(row=1, column=2, value="Nombre").font = Font(bold=True)
    for i, (statut, n) in enumerate(counts.items(), start=2):
        ws2.cell(row=i, column=1, value=statut)
        ws2.cell(row=i, column=2, value=int(n))
        if statut in STATUS_COLORS:
            ws2.cell(row=i, column=1).fill = PatternFill(start_color=STATUS_COLORS[statut], end_color=STATUS_COLORS[statut], fill_type="solid")

    wb.save(output_path)


def plot_progression(df: pd.DataFrame, output_path: Path) -> None:
    fig, ax = plt.subplots(figsize=(10, 6))
    for code, grp in df.groupby("code_indicateur"):
        grp_sorted = grp.sort_values("periode")
        pct = pd.to_numeric(grp_sorted["pct_atteinte"], errors="coerce")
        if pct.notna().any():
            ax.plot(grp_sorted["periode"], pct, marker="o", label=str(code))
    ax.axhline(100, color="green", linestyle="--", linewidth=1, label="Cible (100%)")
    ax.set_ylabel("% d'atteinte de la cible")
    ax.set_xlabel("Période")
    ax.set_title("Progression des indicateurs par période")
    ax.legend(fontsize=8, loc="best")
    fig.tight_layout()
    output_path.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(output_path, dpi=150)
    plt.close(fig)


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--calculated", required=True)
    ap.add_argument("--output-dir", default="output")
    args = ap.parse_args()

    calc_path = Path(args.calculated)
    if not calc_path.exists():
        print(f"Fichier introuvable : {calc_path}", file=sys.stderr)
        return 1

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    run_dir = Path(args.output_dir) / f"dashboard_{timestamp}"
    run_dir.mkdir(parents=True, exist_ok=True)
    logging.basicConfig(level=logging.INFO, format="%(asctime)s [%(levelname)s] %(message)s",
                         handlers=[logging.FileHandler(run_dir / "journal.log", encoding="utf-8"), logging.StreamHandler()])

    df = pd.read_csv(calc_path)
    logger.info("Donnees calculees chargees : %d ligne(s)", len(df))

    dashboard_path = run_dir / "tableau_de_bord.xlsx"
    write_dashboard(df, dashboard_path)
    logger.info("Tableau de bord ecrit : %s", dashboard_path)

    chart_path = run_dir / "progression.png"
    plot_progression(df, chart_path)
    logger.info("Graphique de progression ecrit : %s", chart_path)

    print(f"OK - tableau de bord: {dashboard_path}")
    print(f"OK - graphique: {chart_path}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
