"""
Tests automatises du pipeline survey_intelligence. Executes reellement
avec pytest -- verrouillent les comportements critiques (aucune cible
inventee, classification correcte de severite, detection d'erreur
XLSForm) pour eviter les regressions silencieuses.
"""
import sys
from pathlib import Path

import pandas as pd
import pytest

SCRIPTS_DIR = Path(__file__).resolve().parent.parent / "scripts"
sys.path.insert(0, str(SCRIPTS_DIR))

from indicator_matrix_builder import build_matrix  # noqa: E402
from xlsform_builder import build_xlsform, parse_question_spec, slugify  # noqa: E402
from xlsform_validator import validate_xlsform  # noqa: E402
from survey_quality_checks import (  # noqa: E402
    SEVERITY_MAP,
    build_alert_registry,
    check_consistency_rules,
    check_outliers_iqr,
)


def test_matrix_never_invents_missing_target():
    """Aucune cible absente du CSV source ne doit apparaitre dans la
    matrice -- elle doit rester vide et signalee comme lacune.
    """
    df = pd.DataFrame([{"intitule": "Indicateur test", "definition": "def", "formule": "f",
                         "unite": "u", "source_verification": "sv", "frequence": "freq",
                         "cible": None}])
    matrix_df, gaps = build_matrix(df)
    assert pd.isna(matrix_df.iloc[0]["cible"])
    assert any("cible" in str(g["champs_manquants"]) for g in gaps)


def test_matrix_signals_missing_required_field():
    df = pd.DataFrame([{"intitule": "Indicateur incomplet", "definition": None, "formule": "f",
                         "source_verification": "sv", "frequence": "freq", "cible": "100"}])
    matrix_df, gaps = build_matrix(df)
    gap_for_this = next(g for g in gaps if g["indicateur"] == "Indicateur incomplet")
    assert "definition" in gap_for_this["champs_manquants"]


def test_question_spec_parsing_select_one():
    parsed = parse_question_spec("select_one:Sexe:Masculin,Feminin", "fallback")
    assert parsed["type"] == "select_one"
    assert parsed["choices"] == ["Masculin", "Feminin"]
    assert parsed["reconnu"] is True


def test_question_spec_unrecognized_falls_back_to_text_and_flags():
    parsed = parse_question_spec("type_bizarre:xyz", "Label de secours")
    assert parsed["type"] == "text"
    assert parsed["reconnu"] is False
    assert parsed["label"] == "Label de secours"


def test_slugify_produces_valid_xlsform_name():
    assert slugify("Nom du producteur !") == "nom_du_producteur"
    assert slugify("") == "champ"


def test_xlsform_builder_handles_duplicate_names():
    df = pd.DataFrame([
        {"intitule": "Age", "question_necessaire": "integer:Age"},
        {"intitule": "Age", "question_necessaire": "integer:Age (bis)"},
    ])
    survey_rows, choices_rows, settings, warnings = build_xlsform(df, "Test form")
    names = [r["name"] for r in survey_rows if r["name"] not in ("start", "end")]
    assert len(names) == len(set(names)), "les noms de champs doivent etre uniques"
    assert any("deja utilise" in w for w in warnings)


def test_xlsform_validator_detects_missing_choice_list(tmp_path):
    import openpyxl
    path = tmp_path / "invalide.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "survey"
    ws.append(["type", "name", "label", "required"])
    ws.append(["start", "start", "", ""])
    ws.append(["select_one liste_absente", "q1", "Question", "yes"])
    ws.append(["end", "end", "", ""])
    ws2 = wb.create_sheet("choices")
    ws2.append(["list_name", "name", "label"])
    ws3 = wb.create_sheet("settings")
    ws3.append(["form_title", "form_id"])
    ws3.append(["Titre", "id"])
    wb.save(path)

    result = validate_xlsform(path)
    assert result["valide"] is False
    assert any("liste_absente" in e for e in result["erreurs_bloquantes"])


def test_xlsform_validator_accepts_well_formed_form(tmp_path):
    import openpyxl
    path = tmp_path / "valide.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "survey"
    ws.append(["type", "name", "label", "required"])
    ws.append(["start", "start", "", ""])
    ws.append(["text", "nom", "Nom", "yes"])
    ws.append(["end", "end", "", ""])
    ws2 = wb.create_sheet("choices")
    ws2.append(["list_name", "name", "label"])
    ws3 = wb.create_sheet("settings")
    ws3.append(["form_title", "form_id"])
    ws3.append(["Titre", "id"])
    wb.save(path)

    result = validate_xlsform(path)
    assert result["valide"] is True
    assert result["erreurs_bloquantes"] == []


def test_outlier_detection_never_labeled_as_error():
    """Une valeur hors IQR doit toujours etre classee valeur_inhabituelle,
    jamais erreur_bloquante.
    """
    df = pd.DataFrame({"age": [30, 32, 31, 29, 33, 150]})
    anomalies = check_outliers_iqr(df, ["age"])
    assert len(anomalies) == 1
    severity = SEVERITY_MAP.get(anomalies[0]["type_anomalie"])
    assert severity == "valeur_inhabituelle"
    assert severity != "erreur_bloquante"


def test_consistency_rule_detects_violation():
    df = pd.DataFrame({"age_min": [18, 70], "age_max": [65, 65]})
    rules = pd.DataFrame([{"champ_a": "age_min", "operateur": "<", "champ_b": "age_max"}])
    anomalies = check_consistency_rules(df, rules)
    assert len(anomalies) == 1
    assert anomalies[0]["ligne"] == 1


def test_alert_registry_has_all_required_columns():
    anomalies = [{"type_anomalie": "reponse_manquante", "ligne": 0, "champ": "x", "description": "test"}]
    registry = build_alert_registry(anomalies, "projet_test")
    required_cols = {
        "projet", "id_soumission", "date_detection", "enqueteur", "regle_declenchee",
        "valeur_concernee", "niveau_gravite", "justification", "action_recommandee",
        "responsable_verification", "statut", "historique_corrections",
    }
    assert required_cols.issubset(set(registry.columns))
    assert registry.iloc[0]["niveau_gravite"] == "erreur_bloquante"


if __name__ == "__main__":
    sys.exit(pytest.main([__file__, "-v"]))
