#!/usr/bin/env python3
"""
Controle qualite de soumissions d'enquete (CSV, format export
KoboToolbox). Reutilise les controles de
field2report.cleaning (reponses manquantes, doublons, GPS) plutot que
de les dupliquer, et ajoute : coherence inter-questions (regles
fournies), classification en 3 niveaux (erreur_bloquante /
anomalie_probable / valeur_inhabituelle), et export du registre
d'alertes.

IMPORTANT : une valeur inhabituelle n'est JAMAIS declaree fausse
automatiquement -- elle reste "a verifier", jamais "erreur".

Regles de coherence (--consistency-rules) format CSV :
    champ_a, operateur, champ_b
    operateur in {<, <=, >, >=, ==, !=}
Exemple : age_minimum,<,age_maximum

Usage:
    python survey_quality_checks.py sample_data/soumissions_test.csv \
        --required-fields nom_repondant,commune \
        --numeric-fields age,surface_ha \
        --gps-fields latitude,longitude \
        --dedup-key-fields nom_repondant,commune \
        --consistency-rules sample_data/regles_coherence_test.csv \
        --output-dir output/
"""
from __future__ import annotations

import argparse
import logging
import sys
from datetime import datetime
from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

# Reutilise field2report plutot que de dupliquer sa logique.
FIELD2REPORT_SCRIPTS = Path(__file__).resolve().parent.parent.parent / "field2report" / "scripts"
sys.path.insert(0, str(FIELD2REPORT_SCRIPTS))
from cleaning import (  # noqa: E402
    check_duplicate_respondents,
    check_gps_validity,
    check_missing_required,
)

logger = logging.getLogger("survey_intelligence.survey_quality_checks")

OPERATORS = {
    "<": lambda a, b: a < b, "<=": lambda a, b: a <= b,
    ">": lambda a, b: a > b, ">=": lambda a, b: a >= b,
    "==": lambda a, b: a == b, "!=": lambda a, b: a != b,
}

# Classification des types d'anomalies deja produits par field2report
# vers les 3 niveaux demandes par la mission Survey Intelligence.
SEVERITY_MAP = {
    "reponse_manquante": "erreur_bloquante",
    "champ_obligatoire_absent": "erreur_bloquante",
    "coordonnees_gps_invalides": "erreur_bloquante",
    "doublon_repondant": "anomalie_probable",
    "coordonnees_gps_suspectes_zero": "anomalie_probable",
    "valeur_aberrante_iqr": "valeur_inhabituelle",
    "incoherence_inter_questions": "anomalie_probable",
}

ACTION_RECOMMANDEE = {
    "erreur_bloquante": "Corriger ou rejeter la soumission avant tout usage dans le reporting",
    "anomalie_probable": "Verifier auprès de l'enqueteur ou du repondant ; revisite de terrain si non resolu",
    "valeur_inhabituelle": "Examiner la valeur en contexte -- ne pas rejeter automatiquement",
}


def check_outliers_iqr(df: pd.DataFrame, numeric_fields: list[str]) -> list[dict]:
    """Valeurs hors bornes IQR -- classees valeur_inhabituelle, jamais erreur."""
    anomalies = []
    for fld in numeric_fields:
        if fld not in df.columns:
            continue
        numeric = pd.to_numeric(df[fld], errors="coerce").dropna()
        if len(numeric) < 4:
            continue
        q1, q3 = numeric.quantile([0.25, 0.75])
        iqr = q3 - q1
        low, high = q1 - 1.5 * iqr, q3 + 1.5 * iqr
        mask = (numeric < low) | (numeric > high)
        for idx in numeric[mask].index:
            anomalies.append({
                "type_anomalie": "valeur_aberrante_iqr", "ligne": idx, "champ": fld,
                "description": f"Valeur {numeric.loc[idx].item()!r} hors bornes IQR [{low:.2f}, {high:.2f}] -- a examiner, pas necessairement fausse",
            })
    return anomalies


def check_consistency_rules(df: pd.DataFrame, rules_df: pd.DataFrame) -> list[dict]:
    """Applique des regles de coherence inter-questions EXPLICITEMENT
    fournies -- ne devine jamais de regle metier.
    """
    anomalies = []
    for _, rule in rules_df.iterrows():
        champ_a, op, champ_b = rule.get("champ_a"), rule.get("operateur"), rule.get("champ_b")
        if op not in OPERATORS:
            logger.warning("Regle de coherence ignoree, operateur inconnu : %s", op)
            continue
        if champ_a not in df.columns or champ_b not in df.columns:
            logger.warning("Regle de coherence ignoree, champ(s) absent(s) : %s, %s", champ_a, champ_b)
            continue
        a = pd.to_numeric(df[champ_a], errors="coerce")
        b = pd.to_numeric(df[champ_b], errors="coerce")
        valid = a.notna() & b.notna()
        violated = valid & ~OPERATORS[op](a, b)
        for idx in df[violated].index:
            anomalies.append({
                "type_anomalie": "incoherence_inter_questions", "ligne": idx, "champ": f"{champ_a}/{champ_b}",
                "description": f"Regle violee : {champ_a} ({df.loc[idx, champ_a]}) {op} {champ_b} ({df.loc[idx, champ_b]}) est FAUX",
            })
    return anomalies


def build_alert_registry(anomalies: list[dict], project_id: str) -> pd.DataFrame:
    rows = []
    for a in anomalies:
        severite = SEVERITY_MAP.get(a["type_anomalie"], "anomalie_probable")
        rows.append({
            "projet": project_id,
            "id_soumission": a.get("ligne"),
            "date_detection": datetime.now().strftime("%Y-%m-%d"),
            "enqueteur": a.get("enqueteur", "non_identifie"),
            "regle_declenchee": a["type_anomalie"],
            "valeur_concernee": a.get("champ", ""),
            "niveau_gravite": severite,
            "justification": a["description"],
            "action_recommandee": ACTION_RECOMMANDEE.get(severite, "A examiner"),
            "responsable_verification": "superviseur_a_assigner",
            "statut": "ouverte",
            "historique_corrections": "",
        })
    return pd.DataFrame(rows)


def write_alert_excel_and_csv(df: pd.DataFrame, run_dir: Path) -> tuple[Path, Path]:
    csv_path = run_dir / "registre_alertes.csv"
    df.to_csv(csv_path, index=False)

    excel_path = run_dir / "registre_alertes.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Alertes"
    header_fill = PatternFill(start_color="1F3864", end_color="1F3864", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    severity_fill = {
        "erreur_bloquante": "FFC7CE", "anomalie_probable": "FFEB9C", "valeur_inhabituelle": "D9D9D9",
    }
    cols = list(df.columns)
    for c, h in enumerate(cols, start=1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.fill = header_fill
        cell.font = header_font
    for i, row in enumerate(df.itertuples(index=False), start=2):
        row_dict = dict(zip(cols, row))
        for c, col in enumerate(cols, start=1):
            ws.cell(row=i, column=c, value=row_dict[col])
        sev = row_dict.get("niveau_gravite")
        if sev in severity_fill:
            fill = PatternFill(start_color=severity_fill[sev], end_color=severity_fill[sev], fill_type="solid")
            for c in range(1, len(cols) + 1):
                ws.cell(row=i, column=c).fill = fill
    for c in range(1, len(cols) + 1):
        ws.column_dimensions[get_column_letter(c)].width = 22
    wb.save(excel_path)
    return csv_path, excel_path


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("submissions_file")
    ap.add_argument("--project-id", default="projet_non_specifie")
    ap.add_argument("--required-fields", default="")
    ap.add_argument("--numeric-fields", default="")
    ap.add_argument("--gps-fields", default="", help="2 champs separes par une virgule : lat,lon")
    ap.add_argument("--dedup-key-fields", default="")
    ap.add_argument("--consistency-rules", default=None, help="CSV de regles champ_a,operateur,champ_b")
    ap.add_argument("--output-dir", default="output")
    args = ap.parse_args()

    submissions_path = Path(args.submissions_file)
    if not submissions_path.exists():
        print(f"Fichier introuvable : {submissions_path}", file=sys.stderr)
        return 1

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    run_dir = Path(args.output_dir) / f"quality_{submissions_path.stem}_{timestamp}"
    run_dir.mkdir(parents=True, exist_ok=True)
    logging.basicConfig(level=logging.INFO, format="%(asctime)s [%(levelname)s] %(message)s",
                         handlers=[logging.FileHandler(run_dir / "journal.log", encoding="utf-8"), logging.StreamHandler()])

    df = pd.read_csv(submissions_path)
    logger.info("Soumissions chargees : %d ligne(s)", len(df))

    def _parse_list(v):
        return [x.strip() for x in v.split(",") if x.strip()]

    required_fields = _parse_list(args.required_fields)
    numeric_fields = _parse_list(args.numeric_fields)
    gps_fields = _parse_list(args.gps_fields)
    dedup_fields = _parse_list(args.dedup_key_fields)

    anomalies: list[dict] = []
    anomalies += check_missing_required(df, required_fields)
    if dedup_fields:
        anomalies += check_duplicate_respondents(df, dedup_fields)
    if len(gps_fields) == 2:
        anomalies += check_gps_validity(df, gps_fields[0], gps_fields[1])
    anomalies += check_outliers_iqr(df, numeric_fields)

    if args.consistency_rules:
        rules_path = Path(args.consistency_rules)
        if rules_path.exists():
            rules_df = pd.read_csv(rules_path)
            anomalies += check_consistency_rules(df, rules_df)
        else:
            logger.warning("Fichier de regles de coherence introuvable : %s -- regles ignorees", rules_path)

    for a in anomalies:
        sev = SEVERITY_MAP.get(a["type_anomalie"], "anomalie_probable")
        logger.info("[%s] ligne=%s champ=%s : %s", sev, a.get("ligne"), a.get("champ"), a["description"])

    alert_df = build_alert_registry(anomalies, args.project_id)
    csv_path, excel_path = write_alert_excel_and_csv(alert_df, run_dir)

    n_bloquantes = (alert_df["niveau_gravite"] == "erreur_bloquante").sum() if not alert_df.empty else 0
    n_probables = (alert_df["niveau_gravite"] == "anomalie_probable").sum() if not alert_df.empty else 0
    n_inhabituelles = (alert_df["niveau_gravite"] == "valeur_inhabituelle").sum() if not alert_df.empty else 0

    logger.info(
        "Termine : %d soumission(s) analysee(s), %d erreur(s) bloquante(s), %d anomalie(s) probable(s), %d valeur(s) inhabituelle(s)",
        len(df), n_bloquantes, n_probables, n_inhabituelles,
    )
    print(f"OK - registre alertes (CSV): {csv_path}")
    print(f"OK - registre alertes (Excel): {excel_path}")
    print(f"Resume : {n_bloquantes} erreur(s) bloquante(s), {n_probables} anomalie(s) probable(s), {n_inhabituelles} valeur(s) inhabituelle(s)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
