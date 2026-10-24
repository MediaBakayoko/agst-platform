#!/usr/bin/env python3
"""
CLI EUDR & Forest Monitoring.

IMPORTANT : ce script ne conclut jamais a une infraction ou conformite
EUDR. Il produit des alertes (a examiner) et observations documentees.

Usage:
    python run_eudr_check.py \
        --parcels parcelles.gpkg \
        --forest-baseline foret_ref_2020.tif \
        --forest-recent foret_recent.tif \
        --output-dir output/ \
        --id-field id_parcelle \
        --alert-threshold-moderate-pct 5 \
        --alert-threshold-high-pct 20
"""
from __future__ import annotations

import argparse
import logging
import sys
from datetime import datetime
from pathlib import Path

import geopandas as gpd
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).parent))
from eudr_checks import check_parcel_geolocation, analyze_forest_change, EUDR_REFERENCE_DATE  # noqa: E402


ALERT_COLORS = {
    "aucune": "C6EFCE",
    "faible": "FFEB9C",
    "moderee": "FFC966",
    "elevee": "FFC7CE",
}


def write_report(geoloc_obs, alerts, params, output_path: Path) -> None:
    wb = openpyxl.Workbook()
    header_fill = PatternFill(start_color="1F3864", end_color="1F3864", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)

    ws = wb.active
    ws.title = "Resume"
    ws.cell(row=1, column=1, value="AVERTISSEMENT").font = Font(bold=True, color="C00000", size=12)
    ws.cell(row=2, column=1, value=(
        "Ce rapport contient des ALERTES a examiner par un expert, pas des conclusions "
        "de conformite ou d'infraction EUDR. Une alerte n'equivaut jamais a une "
        "deforestation confirmee."
    ))
    ws.cell(row=4, column=1, value="Date de reference reglementaire EUDR").font = Font(bold=True)
    ws.cell(row=4, column=2, value=str(EUDR_REFERENCE_DATE))
    ws.cell(row=5, column=1, value="Nombre de parcelles analysees").font = Font(bold=True)
    ws.cell(row=5, column=2, value=len(alerts))
    ws.cell(row=6, column=1, value="Observations de geolocalisation").font = Font(bold=True)
    ws.cell(row=6, column=2, value=len(geoloc_obs))

    counts = {lvl: 0 for lvl in ALERT_COLORS}
    for a in alerts:
        counts[a.niveau_alerte] += 1
    ws.cell(row=8, column=1, value="Repartition par niveau d'alerte").font = Font(bold=True, size=12)
    for i, (lvl, n) in enumerate(counts.items(), start=9):
        ws.cell(row=i, column=1, value=lvl)
        cell = ws.cell(row=i, column=2, value=n)
        ws.cell(row=i, column=1).fill = PatternFill(start_color=ALERT_COLORS[lvl], end_color=ALERT_COLORS[lvl], fill_type="solid")

    for c in range(1, 3):
        ws.column_dimensions[get_column_letter(c)].width = 40

    ws2 = wb.create_sheet("Alertes parcelles")
    headers = ["feature_id", "niveau_alerte", "surface_parcelle_ha", "surface_foret_reference_ha",
               "surface_foret_recente_ha", "surface_perte_ha", "pct_perte", "note"]
    for c, h in enumerate(headers, start=1):
        cell = ws2.cell(row=1, column=c, value=h)
        cell.fill = header_fill
        cell.font = header_font
    for i, a in enumerate(alerts, start=2):
        row = [a.feature_id, a.niveau_alerte, a.surface_parcelle_ha, a.surface_foret_reference_ha,
               a.surface_foret_recente_ha, a.surface_perte_ha, a.pct_perte, a.note]
        for c, val in enumerate(row, start=1):
            ws2.cell(row=i, column=c, value=val)
        fill = PatternFill(start_color=ALERT_COLORS[a.niveau_alerte], end_color=ALERT_COLORS[a.niveau_alerte], fill_type="solid")
        for c in range(1, 3):
            ws2.cell(row=i, column=c).fill = fill
    for c in range(1, len(headers) + 1):
        ws2.column_dimensions[get_column_letter(c)].width = 22

    ws3 = wb.create_sheet("Observations geolocalisation")
    headers3 = ["feature_id", "type_observation", "severite", "description"]
    for c, h in enumerate(headers3, start=1):
        cell = ws3.cell(row=1, column=c, value=h)
        cell.fill = header_fill
        cell.font = header_font
    for i, o in enumerate(geoloc_obs, start=2):
        for c, val in enumerate([o.feature_id, o.type_observation, o.severite, o.description], start=1):
            ws3.cell(row=i, column=c, value=val)
    for c in range(1, len(headers3) + 1):
        ws3.column_dimensions[get_column_letter(c)].width = 28

    ws4 = wb.create_sheet("Parametres")
    for i, (k, v) in enumerate(params.items(), start=1):
        ws4.cell(row=i, column=1, value=k)
        ws4.cell(row=i, column=2, value=str(v))
    ws4.column_dimensions["A"].width = 30
    ws4.column_dimensions["B"].width = 50

    wb.save(output_path)


def plot_alert_map(gdf: gpd.GeoDataFrame, alerts, id_field, output_path: Path) -> None:
    alert_by_id = {a.feature_id: a.niveau_alerte for a in alerts}
    gdf = gdf.copy()
    gdf["_niveau_alerte"] = gdf[id_field].map(alert_by_id) if id_field else gdf.index.map(alert_by_id)

    fig, ax = plt.subplots(figsize=(8, 8))
    color_map = {"aucune": "#4CAF50", "faible": "#FFC107", "moderee": "#FF9800", "elevee": "#D32F2F"}
    for lvl, color in color_map.items():
        subset = gdf[gdf["_niveau_alerte"] == lvl]
        if not subset.empty:
            subset.plot(ax=ax, color=color, edgecolor="black", linewidth=0.3, label=lvl)
    ax.legend(title="Niveau d'alerte (a examiner)", fontsize=9)
    ax.set_title("Alertes de changement de couverture forestiere\n(ALERTES, non des conclusions de conformite)", fontsize=10)
    fig.tight_layout()
    output_path.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(output_path, dpi=150)
    plt.close(fig)


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--parcels", required=True)
    ap.add_argument("--forest-baseline", required=True, help="Raster binaire foret (0/1) a la date de reference")
    ap.add_argument("--forest-recent", required=True, help="Raster binaire foret (0/1) recent")
    ap.add_argument("--output-dir", default="output")
    ap.add_argument("--id-field", default=None)
    ap.add_argument("--alert-threshold-moderate-pct", type=float, default=5.0)
    ap.add_argument("--alert-threshold-high-pct", type=float, default=20.0)
    args = ap.parse_args()

    parcels_path = Path(args.parcels)
    if not parcels_path.exists():
        print(f"Fichier parcelles introuvable : {parcels_path}", file=sys.stderr)
        return 1

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    run_dir = Path(args.output_dir) / f"{parcels_path.stem}_{timestamp}"
    run_dir.mkdir(parents=True, exist_ok=True)

    log_path = run_dir / "journal.log"
    logging.basicConfig(
        level=logging.INFO,
        format="%(asctime)s [%(levelname)s] %(message)s",
        handlers=[logging.FileHandler(log_path, encoding="utf-8"), logging.StreamHandler()],
    )
    logger = logging.getLogger("eudr_forest_monitoring.cli")
    logger.info("Demarrage du controle EUDR sur %s", parcels_path)
    logger.info("Date de reference reglementaire : %s", EUDR_REFERENCE_DATE)
    logger.info("Parametres CLI : %s", vars(args))

    gdf = gpd.read_file(parcels_path)
    logger.info("%d parcelle(s) chargee(s), CRS=%s", len(gdf), gdf.crs)

    geoloc_obs = check_parcel_geolocation(gdf, id_field=args.id_field)
    logger.info("Controle geolocalisation : %d observation(s)", len(geoloc_obs))

    alerts = analyze_forest_change(
        gdf, args.forest_baseline, args.forest_recent, id_field=args.id_field,
        seuil_moderee=args.alert_threshold_moderate_pct, seuil_elevee=args.alert_threshold_high_pct,
    )
    for a in alerts:
        logger.info("Parcelle %s : alerte=%s perte=%.2f%%", a.feature_id, a.niveau_alerte, a.pct_perte)

    report_path = run_dir / "rapport_eudr.xlsx"
    write_report(geoloc_obs, alerts, vars(args), report_path)
    logger.info("Rapport Excel ecrit : %s", report_path)

    map_path = run_dir / "carte_alertes.png"
    plot_alert_map(gdf, alerts, args.id_field, map_path)
    logger.info("Carte des alertes ecrite : %s", map_path)

    n_elevee = sum(1 for a in alerts if a.niveau_alerte == "elevee")
    n_moderee = sum(1 for a in alerts if a.niveau_alerte == "moderee")
    logger.info(
        "Termine. %d parcelle(s) analysee(s) : %d alerte(s) elevee(s), %d moderee(s) -- A EXAMINER PAR UN EXPERT.",
        len(alerts), n_elevee, n_moderee,
    )

    print(f"OK - rapport: {report_path}")
    print(f"OK - carte: {map_path}")
    print(f"OK - journal: {log_path}")
    print("RAPPEL : ces resultats sont des alertes a examiner, pas des conclusions de conformite EUDR.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
